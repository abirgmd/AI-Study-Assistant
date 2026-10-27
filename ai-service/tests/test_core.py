"""Phase 1 : extraction, découpage et embeddings (socle partagé)."""

from __future__ import annotations

import io

import numpy as np
import pytest
from docx import Document

from app.core.chunking import chunk_pages, chunk_text, split_sentences
from app.core.embeddings import HashingEmbedder, build_embedder, normalize_rows
from app.core.extraction import Page, clean_text, extract_pages, extract_text
from app.core.text import content_words, normalize, strip_accents
from app.errors import ExtractionError
from tests.fixtures.make_pdf import build_pdf


# --- extraction ---------------------------------------------------------------
def test_txt_utf8(tmp_path):
    path = tmp_path / "cours.txt"
    path.write_text("Les réseaux de neurones apprennent à partir de données étiquetées.", encoding="utf-8")
    assert "réseaux de neurones" in extract_text(path)


def test_txt_latin1_detecte(tmp_path):
    path = tmp_path / "cours.txt"
    path.write_bytes("La modélisation des données relationnelles et les clés étrangères.".encode("latin-1"))
    assert "modélisation" in extract_text(path)


def test_docx_paragraphes_et_tableaux(tmp_path):
    document = Document()
    document.add_paragraph("Introduction au génie logiciel.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Principe"
    table.rows[0].cells[1].text = "SOLID"
    buffer = io.BytesIO()
    document.save(buffer)
    path = tmp_path / "cours.docx"
    path.write_bytes(buffer.getvalue())
    text = extract_text(path)
    assert "génie logiciel" in text and "SOLID" in text


def test_pdf_pages_numerotees(tmp_path):
    path = tmp_path / "cours.pdf"
    path.write_bytes(build_pdf([["Premiere page sur le RAG."], ["Deuxieme page sur FAISS."]]))
    pages = extract_pages(path)
    assert [p.number for p in pages] == [1, 2]
    assert "FAISS" in pages[1].text


def test_pdf_scanne_sans_texte(tmp_path):
    path = tmp_path / "scan.pdf"
    path.write_bytes(build_pdf([[]]))
    with pytest.raises(ExtractionError, match="scanné"):
        extract_pages(path)


def test_pdf_corrompu(tmp_path):
    path = tmp_path / "faux.pdf"
    path.write_bytes(b"ceci n'est pas un pdf")
    with pytest.raises(ExtractionError):
        extract_pages(path)


def test_format_non_pris_en_charge(tmp_path):
    path = tmp_path / "cours.pptx"
    path.write_bytes(b"x")
    with pytest.raises(ExtractionError, match="Format"):
        extract_pages(path)


def test_document_vide(tmp_path):
    path = tmp_path / "vide.txt"
    path.write_text("   \n\n  ", encoding="utf-8")
    with pytest.raises(ExtractionError):
        extract_pages(path)


def test_troncature_max_chars(tmp_path):
    path = tmp_path / "long.txt"
    path.write_text("mot " * 10_000, encoding="utf-8")
    assert len(extract_text(path, max_chars=1000)) <= 1000


def test_clean_text_recolle_les_mots_coupes():
    assert clean_text("déve-\nloppement   logiciel\r\n\n\n\nfin") == "développement logiciel\n\nfin"


# --- normalisation ------------------------------------------------------------
def test_normalisation_francaise():
    assert strip_accents("Génie électrique, ça marche") == "Genie electrique, ca marche"
    assert normalize("L'IA générative !") == "l ia generative"
    assert content_words("Quels sont les 3 avantages du RAG ?") == ["avantages", "rag"]


# --- découpage ----------------------------------------------------------------
def test_split_sentences():
    assert split_sentences("Une phrase. Une autre ! Et une troisième ?") == [
        "Une phrase.", "Une autre !", "Et une troisième ?"]


def test_blocs_respectent_la_taille_et_se_chevauchent():
    text = " ".join(f"Phrase numéro {i} qui parle de tests automatisés." for i in range(60))
    chunks = chunk_text(text, size=300, overlap=80)
    assert len(chunks) > 5
    assert all(len(c.text) <= 300 for c in chunks)
    # Chevauchement : la dernière phrase d'un bloc commence le bloc suivant.
    last_sentence = split_sentences(chunks[0].text)[-1]
    assert chunks[1].text.startswith(last_sentence)
    assert [c.rank for c in chunks] == list(range(len(chunks)))


def test_phrase_trop_longue_coupee_entre_deux_mots():
    text = "mot " * 400
    chunks = chunk_text(text, size=200, overlap=0)
    assert all(len(c.text) <= 200 for c in chunks)
    assert all(not c.text.startswith(" ") for c in chunks)


def test_sans_chevauchement_aucune_repetition():
    text = " ".join(f"Phrase {i}." for i in range(100))
    chunks = chunk_text(text, size=100, overlap=0)
    joined = " ".join(c.text for c in chunks)
    assert joined.split() == text.split()


@pytest.mark.parametrize("size, overlap", [(0, 0), (100, 100), (100, -1)])
def test_parametres_incoherents(size, overlap):
    with pytest.raises(ValueError):
        chunk_text("texte", size=size, overlap=overlap)


def test_blocs_par_page():
    pages = [Page(1, "Page un. " * 5), Page(2, "Page deux. " * 5)]
    chunks = chunk_pages(pages, size=200, overlap=20)
    assert {c.page for c in chunks} == {1, 2}
    assert [c.rank for c in chunks] == list(range(len(chunks)))


def test_texte_vide_aucun_bloc():
    assert chunk_text("", size=100, overlap=10) == []


# --- embeddings ----------------------------------------------------------------
def test_hashing_deterministe_et_normalise():
    embedder = HashingEmbedder()
    a = embedder.encode(["Le protocole TCP", "Les tests unitaires"])
    b = HashingEmbedder().encode(["Le protocole TCP", "Les tests unitaires"])
    assert a.dtype == np.float32 and a.shape == (2, embedder.dim)
    np.testing.assert_allclose(a, b)
    np.testing.assert_allclose(np.linalg.norm(a, axis=1), 1.0, rtol=1e-5)


def test_hashing_rapproche_les_textes_proches():
    embedder = HashingEmbedder()
    q, proche, loin = embedder.encode([
        "Qu'est-ce que la couverture de code ?",
        "La couverture de code mesure le pourcentage de code exécuté par les tests.",
        "Le protocole UDP ne garantit pas la livraison des paquets.",
    ])
    assert float(q @ proche) > float(q @ loin)


def test_build_embedder_unique_par_configuration():
    assert build_embedder("hashing", "x") is build_embedder("hashing", "x")
    with pytest.raises(ValueError):
        build_embedder("inconnu", "x")


def test_normalize_rows_ligne_nulle():
    result = normalize_rows(np.zeros((1, 4)))
    assert not np.isnan(result).any()
