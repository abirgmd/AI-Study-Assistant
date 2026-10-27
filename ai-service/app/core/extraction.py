"""Extraction du texte des cours importés (PDF, DOCX, TXT).

Module partagé : utilisé par le RAG (module 3), la classification (module 5) et
la recherche sémantique (module 2).
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from pathlib import Path

from charset_normalizer import from_bytes

from app.errors import ExtractionError

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
_HYPHEN_BREAK = re.compile(r"(\w)-\n(\w)")
_SPACES = re.compile(r"[ \t ]+")
_MANY_NEWLINES = re.compile(r"\n{3,}")


@dataclass(frozen=True)
class Page:
    """Bloc de texte extrait, avec son numéro de page (None pour TXT et DOCX)."""

    number: int | None
    text: str


def clean_text(text: str) -> str:
    """Normalise le texte extrait sans en changer le sens.

    Recolle les mots coupés en fin de ligne, réduit les espaces multiples,
    supprime les caractères de contrôle et limite les sauts de ligne.

    Args:
        text: Texte brut.

    Returns:
        Texte nettoyé.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = "".join(ch for ch in text if ch == "\n" or ch == "\t" or ch.isprintable())
    text = _HYPHEN_BREAK.sub(r"\1\2", text)
    text = _SPACES.sub(" ", text)
    lines = [line.strip() for line in text.split("\n")]
    text = "\n".join(lines)
    return _MANY_NEWLINES.sub("\n\n", text).strip()


def _extract_pdf(path: Path) -> list[Page]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
        pages = [Page(i + 1, clean_text(page.extract_text() or "")) for i, page in enumerate(reader.pages)]
    except (PdfReadError, ValueError, KeyError) as exc:
        logger.warning("PDF illisible : %s", exc)
        raise ExtractionError("Le PDF est illisible ou protégé.") from exc
    pages = [p for p in pages if p.text]
    if not pages:
        raise ExtractionError(
            "Ce PDF ne contient pas de texte sélectionnable (document scanné ?). "
            "Importez une version texte du cours."
        )
    return pages


def _extract_docx(path: Path) -> list[Page]:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = Document(str(path))
    except (PackageNotFoundError, KeyError, ValueError) as exc:
        raise ExtractionError("Le document DOCX est invalide.") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return [Page(None, clean_text("\n".join(parts)))]


def _extract_txt(path: Path) -> list[Page]:
    raw = path.read_bytes()
    best = from_bytes(raw).best()
    text = str(best) if best is not None else raw.decode("utf-8", errors="replace")
    return [Page(None, clean_text(text))]


def extract_pages(path: Path, max_chars: int = 400_000) -> list[Page]:
    """Extrait le texte d'un document, page par page pour les PDF.

    Args:
        path: Chemin du fichier (déjà validé par l'appelant).
        max_chars: Nombre maximal de caractères conservés (protection mémoire).

    Returns:
        Liste non vide de pages contenant du texte.

    Raises:
        ExtractionError: Format non pris en charge, document vide ou illisible.
    """
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ExtractionError("Format non pris en charge (PDF, DOCX ou TXT uniquement).")
    if extension == ".pdf":
        pages = _extract_pdf(path)
    elif extension == ".docx":
        pages = _extract_docx(path)
    else:
        pages = _extract_txt(path)

    kept: list[Page] = []
    total = 0
    for page in pages:
        if not page.text:
            continue
        remaining = max_chars - total
        if remaining <= 0:
            logger.info("Document tronqué à %d caractères", max_chars)
            break
        kept.append(Page(page.number, page.text[:remaining]))
        total += min(len(page.text), remaining)
    if not kept:
        raise ExtractionError("Le document ne contient aucun texte exploitable.")
    return kept


def extract_text(path: Path, max_chars: int = 400_000) -> str:
    """Extrait tout le texte d'un document en une seule chaîne.

    Args:
        path: Chemin du fichier.
        max_chars: Nombre maximal de caractères conservés.

    Returns:
        Texte complet du document.
    """
    return "\n\n".join(page.text for page in extract_pages(path, max_chars))
